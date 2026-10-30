import io
from datetime import date
from decimal import Decimal
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook

from datumoj.rikolto.core.contracts import ParseError
from datumoj.rikolto.core.models import parse_locator
from datumoj.rikolto.parsers.mhlw_ndb import MEMBER, MhlwNdbParser
from datumoj.rikolto.testing.contract_suites import ParserSuite
from datumoj.rikolto.testing.samples import ParseSample, make_snapshot

ROOT = Path(__file__).parents[2]
XLSX = (ROOT / 'tests/fixtures/ndb/ndb11_kenshin_shitsumon01_pref_sex_age.xlsx').read_bytes()
URL = 'https://www.mhlw.go.jp/content/12400000/001711941.zip'


def archive(payload=XLSX, member=MEMBER, legacy=False):
    target = io.BytesIO()
    name = 'a' * len(member.encode('cp932')) if legacy else member
    with ZipFile(target, 'w', ZIP_DEFLATED) as z:
        z.writestr(name, payload)
        z.writestr('other.xlsx', XLSX)
    data = target.getvalue()
    return data.replace(name.encode(), member.encode('cp932')) if legacy else data


DATA = archive()


def sample(data=DATA):
    return ParseSample(make_snapshot('jp-mhlw', URL, data, content_type='text/plain; charset=ascii'),
                       data, {'member': MEMBER})


def parse(data=DATA):
    s = sample(data)
    return MhlwNdbParser().parse(s.snapshot, s.data, s.meta)


class TestMhlwNdbParser(ParserSuite):
    @pytest.fixture
    def parser(self):
        return MhlwNdbParser()

    @pytest.fixture
    def samples(self):
        return [sample()]

    @pytest.fixture
    def bad_samples(self):
        return [b'not a ZIP', DATA[:100], archive(b'not an XLSX')]


def test_values_dimensions_units_and_period_match_recorded_workbook():
    result = parse()
    assert len(result.observations) == 48 * 2 * 2 * 8
    book = load_workbook(io.BytesIO(XLSX), read_only=True, data_only=True)
    try:
        sheet = book['質問項目']
        obs = next(o for o in result.observations if o.area_label == '北海道'
                   and o.dims == {'sex': '男', 'age': '40～44歳', 'answer': 'はい'})
        assert obs.value == Decimal(sheet['C6'].value) == 4688
        assert obs.unit == sheet['C5'].value == '人数'
        assert parse_locator(obs.locator) == {'member': MEMBER, 'sheet': '質問項目', 'cell': 'C6'}
        cells = {c.coordinate: c.value for row in sheet for c in row if c.value is not None}
        for o in result.observations:
            assert o.value == Decimal(cells[parse_locator(o.locator)['cell']])
            assert o.survey == '第11回NDBオープンデータ 特定健診'
            assert (o.time_start, o.time_end, o.time_label) == (date(2023, 4, 1), date(2024, 3, 31), '2023年度')
        assert any(o.area_label == '都道府県判別不可' for o in result.observations)
        assert any(o.dims['age'] == '中計' for o in result.observations)
    finally:
        book.close()


def modified(tmp_path, updates):
    book = load_workbook(io.BytesIO(XLSX))
    for cell, value in updates.items():
        for merged in list(book.active.merged_cells.ranges):
            if cell in merged:
                book.active.unmerge_cells(str(merged))
        book.active[cell] = value
    path = tmp_path / 'ndb.synthetic.xlsx'
    book.save(path)
    book.close()
    return archive(path.read_bytes())


@pytest.mark.parametrize('mark', ['-', '…', 'x', '秘匿', None])
def test_suppression_and_missing_counts_never_become_zero(tmp_path, mark):
    # This derived fixture has no suppressed cells; change only a synthetic copy.
    result = parse(modified(tmp_path, {'C6': mark}))
    obs = next(o for o in result.observations if parse_locator(o.locator)['cell'] == 'C6')
    assert obs.value is None and obs.note
    if mark == '-':
        assert obs.note == '-; 遮蔽（集计单位小于 10 或防反推）'
    elif mark is not None:
        assert obs.note == mark


@pytest.mark.parametrize('legacy', [False, True])
def test_exact_fullwidth_member_name_from_zip(tmp_path, legacy):
    path = tmp_path / 'ndb.synthetic.zip'
    path.write_bytes(archive(legacy=legacy))
    assert len(parse(path.read_bytes()).observations) == 1536
    # A different item is present in the archive; it must not be used as a fallback.
    path.write_bytes(archive(member=MEMBER.replace('１', '1'), legacy=legacy))
    with pytest.raises(ParseError, match=MEMBER.replace('（', r'\（').replace('）', r'\）')):
        parse(path.read_bytes())


@pytest.mark.parametrize('updates', [
    {'A1': '別の質問票：2023年度'}, {'C3': '合計'}, {'C4': '40-44'}, {'C5': '千人'},
    {'B6': '不明'}, {'A8': '北海道'}, {'A7': '北海道'}, {'C6': '=1+2'}, {'C6': -1},
    {'C6': 'not a count'}, {'A102': 'extra data'},
])
def test_invalid_layout_or_counts_fail_atomically(tmp_path, updates):
    with pytest.raises(ParseError):
        parse(modified(tmp_path, updates))


def test_wrong_metadata_is_rejected():
    s = sample()
    with pytest.raises(ParseError, match='requires ZIP member'):
        MhlwNdbParser().parse(s.snapshot, s.data, {'member': MEMBER.replace('１', '1')})
