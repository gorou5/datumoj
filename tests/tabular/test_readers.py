import csv
import io
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

import pytest
from openpyxl import Workbook

from datumoj.rikolto import tabular
from datumoj.rikolto.core.models import parse_locator
from datumoj.rikolto.tabular import ArchiveLimitError, detect_header, read_csv, read_xlsx, read_zip, sniff

FIXTURES = Path(__file__).parents[1] / "fixtures" / "tabular"


def fixture(name):
    return (FIXTURES / name).read_bytes()


def archive(entries, compression=0):
    output = io.BytesIO()
    with ZipFile(output, "w", compression=compression) as zipped:
        for name, data in entries:
            zipped.writestr(name, data)
    return output.getvalue()


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "表;=1%"
    sheet.append(["地域", "値", "未計算", "キャッシュ"])
    sheet.append(["甲市", " 7,500", "=1+2", "=40+2"])
    workbook.create_sheet("第二表").append(["x", "秘匿"])
    output = io.BytesIO()
    workbook.save(output)
    with ZipFile(output) as zipped:
        entries = [(info.filename, zipped.read(info)) for info in zipped.infolist()]
    return archive([(name, data.replace(b"<f>40+2</f><v></v>", b"<f>40+2</f><v>42</v>")) for name, data in entries])


@pytest.mark.parametrize(
    ("name", "expected"),
    [
        ("estat_ecensus2016_t2_osaka.csv", "csv"),
        ("estat_eisei_gyosei_r6_t10.csv", "csv"),
        ("bodik_mukou_population.xlsx", "xlsx"),
        ("ksj_s12_mixed_encoding.synthetic.zip", "zip"),
        ("zip_not_xlsx_content_types.synthetic.zip", "zip"),
    ],
)
def test_sniff_only_uses_bytes(name, expected):
    assert sniff(fixture(name)) == expected


def test_sniff_xls_unsupported_binary_and_incomplete_xlsx():
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1rest") == "xls"
    assert sniff(b"\x00\x01\x02") == "unsupported"
    assert sniff(b"<xml/>") == "unsupported"
    assert sniff(b'{"a":1}') == "unsupported"
    assert sniff(archive([("xl/workbook.xml", b"not a workbook")])) == "zip"
    assert sniff(archive([])) == "zip"
    with pytest.raises(BadZipFile):
        sniff(b"PK\x03\x04broken")


@pytest.mark.parametrize("delimiter", [",", "\t", ";", "|"])
@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-16", "cp932"])
def test_csv_preserves_strings_and_record_positions(delimiter, encoding):
    output = io.StringIO(newline="")
    source = [["地域", "数値", "注記"], [" 甲市 ", "001", "x"], ["乙\n町", " 7,500", ""]]
    csv.writer(output, delimiter=delimiter).writerows(source)
    data = output.getvalue().encode(encoding)
    assert sniff(data) == "csv"
    table = read_csv(data, member="資料;=%.bin")
    assert [[c.raw for c in row] for row in table.rows] == source
    assert table.sheet is None
    assert table.member == "資料;=%.bin"
    cell = table.rows[2][1]
    assert (cell.row, cell.column) == (3, 2)
    assert parse_locator(cell.locator) == {"member": table.member, "cell": "B3", "row": "3"}


@pytest.mark.parametrize("name", ["estat_ecensus2016_t2_osaka.csv", "estat_eisei_gyosei_r6_t10.csv"])
def test_real_cp932_csv_matches_every_original_field(name):
    data = fixture(name)
    with pytest.raises(UnicodeDecodeError):
        data.decode("utf-8")
    expected = list(csv.reader(io.StringIO(data.decode("cp932"))))
    table = read_csv(data)
    assert [[cell.raw for cell in row] for row in table.rows] == expected
    assert any(cell.raw == "-" for row in table.rows for cell in row)


def test_csv_ragged_preamble_empty_and_single_column():
    table = read_csv("題名\n注：説明\n地域;総数;男\n甲市;10;4\n".encode())
    assert [[c.raw for c in row] for row in table.rows] == [
        ["題名"],
        ["注：説明"],
        ["地域", "総数", "男"],
        ["甲市", "10", "4"],
    ]
    assert read_csv(b"").rows == ()
    assert read_csv("資料\n甲市\n".encode()).rows[1][0].raw == "甲市"
    with pytest.raises(csv.Error):
        read_csv(b'a,b\n"unterminated')


def test_xlsx_read_only_cached_formulas_sheets_and_locators():
    with patch.object(tabular, "load_workbook", wraps=tabular.load_workbook) as loader:
        tables = read_xlsx(workbook_bytes(), member="偽.csv")
    assert loader.call_args.kwargs == {"read_only": True, "data_only": True, "keep_links": False}
    assert [t.sheet for t in tables] == ["表;=1%", "第二表"]
    assert [c.raw for c in tables[0].rows[1]] == ["甲市", " 7,500", "", "42"]
    assert parse_locator(tables[0].rows[1][3].locator) == {"member": "偽.csv", "sheet": "表;=1%", "cell": "D2"}
    assert [c.raw for c in tables[1].rows[0]] == ["x", "秘匿"]


def test_xlsx_ignores_dimension_hint_and_limits_expanded_cells(monkeypatch):
    data = workbook_bytes()
    with ZipFile(io.BytesIO(data)) as zipped:
        entries = [(info.filename, zipped.read(info).replace(b"A1:D2", b"A1:XFD1048576")) for info in zipped.infolist()]
    assert len(read_xlsx(archive(entries))[0].rows) == 2
    monkeypatch.setattr(tabular, "MAX_CELLS", 3)
    with pytest.raises(ValueError, match="cell count"):
        read_xlsx(data)


def test_real_bodik_date_and_one_line_header():
    (table,) = read_xlsx(fixture("bodik_mukou_population.xlsx"))
    assert table.rows[1][3].raw == "2024-08-31 00:00:00"
    assert table.rows[1][0].raw == "262081"
    assert detect_header(table).header_rows == (1,)


@pytest.mark.parametrize(
    ("name", "headers", "start", "notes", "column", "expected"),
    [
        (
            "estat_eisei_gyosei_r6_t10.csv",
            (6, 7, 8, 9),
            10,
            (4, 5),
            5,
            "理容所 / 重複開設（再掲） / 使用確認件数（年度中） / 件",
        ),
        ("estat_ecensus2016_t2_osaka.csv", (12, 13), 14, (6, 7), 8, "総数(経営組織) / 従業者数(人) (注2)"),
        ("estat_chiho_kyoikuhi_h28_t7-1.xlsx", (7, 8), 9, (5,), 3, "地方債・寄附金以外の公費 / 国庫補助金"),
        (
            "estat_hogosha_r2_t1-1.xlsx",
            (5, 6, 7, 8),
            9,
            (3,),
            5,
            "被保護実世帯数 / 現に保護を受けた世帯 / 総数 / （世帯）",
        ),
    ],
)
def test_real_multiline_headers_and_notes(name, headers, start, notes, column, expected):
    data = fixture(name)
    table = read_csv(data) if sniff(data) == "csv" else read_xlsx(data)[0]
    result = detect_header(table)
    assert result.header_rows == headers
    assert result.data_start == start
    assert result.note_rows == notes
    assert result.columns[column] == expected
    if name.startswith("estat_chiho"):
        assert table.rows[9][0].raw.endswith(" ")
        assert {"-", "…"} <= {cell.raw for row in table.rows for cell in row}


def test_true_merges_notes_and_missing_markers_preserve_raw_cells():
    (table,) = read_xlsx(fixture("stat_merged_header_notes.synthetic.xlsx"))
    original = table.rows
    assert "B3:E3" in table.merged_ranges
    assert table.rows[2][2].raw == ""
    assert table.rows[3][0].raw == ""
    result = detect_header(table)
    assert result.header_rows == (3, 4, 5)
    assert result.data_start == 6
    assert result.note_rows == (14, 15, 16, 17)
    assert result.columns[0] == "地域"
    assert result.columns[3] == "事業所数 / 経営組織別 / 法人"
    assert result.columns[7] == "従業者数 / 男女別 / 女"
    assert {"-", "－", "…", "x", "X", "秘匿", "△0.8", "▲2.5", "*", " 7,500", "0.0"} <= {
        cell.raw for row in table.rows for cell in row
    }
    assert table.rows == original
    assert table.rows[13][0].raw.startswith("注：")
    assert table.rows[15][0].raw.startswith("資料：")
    assert parse_locator(table.rows[8][3].locator) == {"sheet": "第3表", "cell": "D9"}


def test_header_empty_and_no_numeric_data():
    assert detect_header(read_csv(b"")).data_start is None
    assert detect_header(read_csv("題名\n地域,説明\n甲市,文章\n".encode())).header_rows == ()


def test_zip_content_types_decoy_and_mixed_names_and_content():
    (table,) = read_zip(fixture("zip_not_xlsx_content_types.synthetic.zip"))
    assert table.member == "data/sample.csv"
    tables = read_zip(fixture("ksj_s12_mixed_encoding.synthetic.zip"))
    names = [f"S12-25_GML/{folder}/駅別乗降客数_説明.txt" for folder in ("Shift-JIS", "UTF-8")]
    selected = [table for table in tables if table.member in names]
    assert {table.member for table in selected} == set(names)
    with ZipFile(io.BytesIO(fixture("ksj_s12_mixed_encoding.synthetic.zip"))) as zipped:
        for table, encoding in zip(selected, ("cp932", "utf-8"), strict=True):
            info = next(
                info
                for info in zipped.infolist()
                if info.filename.endswith(".txt") and ((info.flag_bits & 0x800) != 0) == (encoding == "utf-8")
            )
            assert [[cell.raw for cell in row] for row in table.rows] == list(
                csv.reader(io.StringIO(zipped.read(info).decode(encoding)))
            )
    assert "Shift_JIS(cp932)" in selected[0].rows[0][0].raw
    assert "UTF-8。ZIP" in selected[1].rows[0][0].raw
    for table in selected:
        assert parse_locator(table.rows[0][0].locator)["member"] == table.member


def test_zip_recurses_by_content_even_with_misleading_names():
    csv_bytes = "地域,数値\n甲市,秘匿\n".encode("cp932")
    nested = archive([("資料.xlsx", csv_bytes), ("ブック.csv", workbook_bytes())])
    tables = read_zip(archive([("内側.txt", nested), ("empty/", b"")]))
    assert [t.member for t in tables] == ["内側.txt!/資料.xlsx", "内側.txt!/ブック.csv", "内側.txt!/ブック.csv"]
    assert tables[0].rows[1][1].raw == "秘匿"
    assert parse_locator(tables[1].rows[1][0].locator)["member"] == "内側.txt!/ブック.csv"
    with pytest.raises(NotImplementedError, match="unsupported XLS"):
        read_zip(archive([("old.csv", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")]))


@pytest.mark.parametrize(
    "name",
    [
        "zip_bomb_high_ratio.synthetic.zip",
        "zip_bomb_many_members.synthetic.zip",
        "zip_bomb_declared_size.synthetic.zip",
    ],
)
@pytest.mark.parametrize("reader", [read_zip, read_xlsx])
def test_all_fixture_bombs_rejected_before_opening_any_member(name, reader):
    with (
        patch.object(ZipFile, "open", side_effect=AssertionError("must reject before inflation")),
        pytest.raises(ArchiveLimitError),
    ):
        reader(fixture(name))


def test_ratio_limit_below_size_limit():
    data = archive([("small.csv", b"0" * 65536)], ZIP_DEFLATED)
    with pytest.raises(ArchiveLimitError, match="ratio"):
        read_zip(data)


def test_nested_archives_share_member_size_and_depth_budgets(monkeypatch):
    inner = archive([("a", b"1"), ("b", b"2")])
    outer = archive([("nested", inner), ("c", b"3")])
    monkeypatch.setattr(tabular, "MAX_MEMBERS", 3)
    with pytest.raises(ArchiveLimitError, match="member count"):
        read_zip(outer)
    monkeypatch.setattr(tabular, "MAX_MEMBERS", 1000)
    monkeypatch.setattr(tabular, "MAX_UNCOMPRESSED", len(inner) + 1)
    with pytest.raises(ArchiveLimitError, match="uncompressed size"):
        read_zip(outer)
    monkeypatch.setattr(tabular, "MAX_UNCOMPRESSED", 64 * 1024 * 1024)
    monkeypatch.setattr(tabular, "MAX_DEPTH", 1)
    with pytest.raises(ArchiveLimitError, match="nesting depth"):
        read_zip(outer)


def test_nested_xlsx_uses_archive_budget(monkeypatch):
    data = workbook_bytes()
    monkeypatch.setattr(tabular, "MAX_MEMBERS", 2)
    with pytest.raises(ArchiveLimitError, match="member count"):
        read_zip(archive([("book", data)]))


@pytest.mark.parametrize("name", ["../a.csv", "/a.csv", "a/../../b", "a\\b.csv"])
def test_zip_unsafe_paths_rejected(name):
    with pytest.raises(ValueError, match="unsafe ZIP member"):
        read_zip(archive([(name, b"a,b\n1,2")]))


def test_zip_crc_failure_is_not_silently_dropped():
    data = archive([("table", b"a,b\n1,2")]).replace(b"a,b\n1,2", b"a,b\n1,3")
    with pytest.raises(BadZipFile):
        read_zip(data)
