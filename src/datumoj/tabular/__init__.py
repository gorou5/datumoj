"""Offline, content-sniffed tables. Coordinates are one-based; raw cells are never filled."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import PurePosixPath
from typing import Literal
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter, range_boundaries

from datumoj.core.models import format_locator
from datumoj.text import decode_bytes, parse_number

__all__ = [
    "ArchiveLimitError",
    "HeaderDetection",
    "RawCell",
    "RawTable",
    "detect_header",
    "read_csv",
    "read_xlsx",
    "read_zip",
    "sniff",
]

MAX_MEMBERS = 1000
MAX_UNCOMPRESSED = 64 * 1024 * 1024
MAX_RATIO = 200
MAX_DEPTH = 8
MAX_CELLS = 1_000_000
_SHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


@dataclass(frozen=True)
class RawCell:
    row: int
    column: int
    raw: str
    locator: str


@dataclass(frozen=True)
class RawTable:
    rows: tuple[tuple[RawCell, ...], ...]
    sheet: str | None = None
    member: str | None = None
    merged_ranges: tuple[str, ...] = ()


@dataclass(frozen=True)
class HeaderDetection:
    header_rows: tuple[int, ...]
    columns: tuple[str, ...]
    note_rows: tuple[int, ...]
    data_start: int | None


class ArchiveLimitError(ValueError):
    """An archive exceeds the shared expansion budget."""


@dataclass
class _Budget:
    members: int = 0
    size: int = 0

    def check(self, archive: ZipFile) -> None:
        entries = archive.infolist()
        self.members += len(entries)
        self.size += sum(entry.file_size for entry in entries)
        if self.members > MAX_MEMBERS:
            raise ArchiveLimitError("ZIP member count exceeds limit")
        if self.size > MAX_UNCOMPRESSED:
            raise ArchiveLimitError("ZIP uncompressed size exceeds limit")
        if any(entry.file_size > MAX_RATIO * max(entry.compress_size, 1) for entry in entries):
            raise ArchiveLimitError("ZIP compression ratio exceeds limit")


def sniff(data: bytes) -> Literal["csv", "xlsx", "xls", "zip", "unsupported"]:
    """Inspect bytes and central-directory names, without decompressing ZIP members."""
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls"
    if data.startswith((b"PK\x03\x04", b"PK\x05\x06", b"PK\x07\x08")):
        with ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
        return "xlsx" if {"[Content_Types].xml", "xl/workbook.xml"} <= names else "zip"
    try:
        text, _, _ = decode_bytes(data)
    except UnicodeDecodeError:
        return "unsupported"
    if any(ord(char) < 32 and char not in "\r\n\t" for char in text):
        return "unsupported"
    if text.lstrip().startswith(("<", "{", "[")):
        return "unsupported"
    return "csv"


def _table(rows, sheet=None, member=None, merged_ranges=()) -> RawTable:
    result = []
    count = 0
    for row_number, row in enumerate(rows, 1):
        cells = []
        for column, value in enumerate(row, 1):
            count += 1
            if count > MAX_CELLS:
                raise ValueError("table cell count exceeds limit")
            parts = {"member": member} if member is not None else {}
            if sheet is not None:
                parts["sheet"] = sheet
            parts["cell"] = f"{get_column_letter(column)}{row_number}"
            if sheet is None:
                parts["row"] = row_number
            cells.append(RawCell(row_number, column, "" if value is None else str(value), format_locator(parts)))
        result.append(tuple(cells))
    return RawTable(tuple(result), sheet, member, merged_ranges)


def read_csv(data: bytes, *, member: str | None = None) -> RawTable:
    """Decode by content; keep whitespace, quoted newlines, empty fields and notes."""
    text, _, _ = decode_bytes(data)
    try:
        dialect = csv.Sniffer().sniff(text[:65536], delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        # Government files often have ragged title/note rows before the table.
        scores = {}
        for candidate in (",", "\t", ";", "|"):
            lengths = [len(row) for row in csv.reader(io.StringIO(text), delimiter=candidate) if len(row) > 1]
            scores[candidate] = sum(lengths)
        delimiter = max(scores, key=scores.get)
    return _table(csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True), member=member)


def read_xlsx(data: bytes, *, member: str | None = None) -> list[RawTable]:
    return _read_xlsx(data, member, _Budget())


def _read_xlsx(data: bytes, member: str | None, budget: _Budget) -> list[RawTable]:
    with ZipFile(io.BytesIO(data)) as archive:
        budget.check(archive)
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
        try:
            tables = []
            for sheet in workbook:
                # Read-only openpyxl omits merges; inspect only their XML metadata.
                merges = []
                with archive.open(sheet._worksheet_path) as source:
                    for _, element in ET.iterparse(source, events=("end",)):
                        if element.tag == _SHEET_NS + "mergeCell":
                            merges.append(element.attrib["ref"])
                        element.clear()
                # Ignore untrusted dimension hints (including artificially huge used ranges).
                sheet.reset_dimensions()
                tables.append(_table(sheet.iter_rows(values_only=True), sheet.title, member, tuple(merges)))
            return tables
        finally:
            workbook.close()


def _member_name(info) -> str:
    name = info.filename
    if not info.flag_bits & 0x800:
        try:
            name = name.encode("cp437").decode("cp932")
        except UnicodeDecodeError:
            pass
    # No extraction takes place, but reject unsafe or ambiguous archive paths.
    path = PurePosixPath(name)
    if "\x00" in info.orig_filename or "\\" in name or path.is_absolute() or ".." in path.parts:
        raise ValueError("unsafe ZIP member path")
    return name


def read_zip(data: bytes) -> list[RawTable]:
    """Read supported members recursively in memory; legacy XLS raises unsupported."""
    return _read_zip(data, None, _Budget(), 0)


def _read_zip(data: bytes, parent: str | None, budget: _Budget, depth: int) -> list[RawTable]:
    if depth >= MAX_DEPTH:
        raise ArchiveLimitError("ZIP nesting depth exceeds limit")
    tables = []
    with ZipFile(io.BytesIO(data)) as archive:
        budget.check(archive)
        for info in archive.infolist():
            name = _member_name(info)
            if info.is_dir():
                continue
            member = f"{parent}!/{name}" if parent is not None else name
            # Read one byte past the declaration and let zipfile validate CRC/length.
            with archive.open(info) as source:
                payload = source.read(min(info.file_size, MAX_UNCOMPRESSED) + 1)
            if len(payload) != info.file_size:
                raise ValueError("ZIP member size mismatch")
            kind = sniff(payload)
            if kind == "zip":
                tables.extend(_read_zip(payload, member, budget, depth + 1))
            elif kind == "xlsx":
                tables.extend(_read_xlsx(payload, member, budget))
            elif kind == "csv":
                tables.append(read_csv(payload, member=member))
            elif kind == "xls":
                raise NotImplementedError(f"unsupported XLS member: {member}")
    return tables


_NOTE = re.compile(r"^(?:[（(]?注(?:[:：)）0-9０-９]|\s|$)|資料[:：]|出典[:：])")
_CONTINUATION = re.compile(r"^[0-9０-９]+[)）]")
_MARKERS = {"-", "－", "…", "x", "X", "秘匿", "*", "***", "・"}


def detect_header(table: RawTable) -> HeaderDetection:
    """Infer up to four header rows without changing raw values or data rows.

    shortcut: numeric column headings can resemble data, use layout-specific validation in statistical converters.
    """
    values = [[cell.raw for cell in row] for row in table.rows]
    width = max((i + 1 for row in values for i, value in enumerate(row) if value.strip()), default=0)
    values = [row[:width] + [""] * max(0, width - len(row)) for row in values]
    # e-Stat CSV sometimes prefixes every record, including titles, with its row number.
    if values and all(row and row[0] == str(index) for index, row in enumerate(values, 1)):
        for row in values:
            row[0] = ""
    for merged in table.merged_ranges:
        left, top, right, bottom = range_boundaries(merged)
        if top > len(values) or left > width:
            continue
        anchor = values[top - 1][left - 1]
        for row in range(top - 1, min(bottom, len(values))):
            for column in range(left - 1, min(right, width)):
                values[row][column] = anchor

    notes = []
    in_note = False
    for index, row in enumerate(values):
        nonempty = [value.strip() for value in row if value.strip()]
        first = nonempty[0] if nonempty else ""
        indented = any(value and value != value.lstrip() for value in row)
        if _NOTE.match(first) or (in_note and len(nonempty) == 1 and (_CONTINUATION.match(first) or indented)):
            notes.append(index + 1)
            in_note = True
        elif nonempty:
            in_note = False

    data_index = None
    for index, row in enumerate(values):
        if index + 1 in notes:
            continue
        numeric = [parse_number(value)[0] is not None or value.strip() in _MARKERS for value in row]
        labels = [i for i, value in enumerate(row) if value.strip() and not numeric[i]]
        if labels and any(numeric[1:]):
            data_index = index
            break

    headers = []
    if data_index is not None:
        for index in range(data_index - 1, max(-1, data_index - 5), -1):
            row = values[index]
            nonempty = [value.strip() for value in row if value.strip()]
            if index + 1 in notes or not nonempty or any(parse_number(v)[0] is not None for v in nonempty):
                break
            if len(nonempty) == 1 and (row[0].strip() or "単位" in nonempty[0]):
                break
            headers.insert(0, index + 1)

    header_values = [values[index - 1][:] for index in headers]
    for index, row in enumerate(header_values[:-1]):
        previous = ""
        for column, value in enumerate(row):
            if value.strip():
                previous = value
            elif any(lower[column].strip() for lower in header_values[index + 1 :]):
                row[column] = previous
    columns = []
    for column in range(width):
        parts = []
        for row in header_values:
            value = row[column].strip()
            if value and (not parts or parts[-1] != value):
                parts.append(value)
        columns.append(" / ".join(parts))
    return HeaderDetection(
        tuple(headers), tuple(columns), tuple(notes), data_index + 1 if data_index is not None else None
    )
