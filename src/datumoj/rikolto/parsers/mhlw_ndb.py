"""NDB 11, questionnaire item 1: preserve every published cell and its locator."""

import io
from collections.abc import Mapping
from datetime import date
from decimal import Decimal

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from pydantic import JsonValue

from datumoj.rikolto.core.contracts import ParseError, Parser, ParseResult
from datumoj.rikolto.core.models import Chunk, Document, Observation, Snapshot, TimeGranularity, format_locator
from datumoj.rikolto.parsers._zip import read_zip

MEMBER = "08_特定健診_問診項目/標準的な質問票（質問項目１）　都道府県別性年齢階級別分布.xlsx"
SURVEY = "第11回NDBオープンデータ 特定健診"
INDICATOR = "現在、血圧を下げる薬を使用しているか"
AGES = ("40～44歳", "45～49歳", "50～54歳", "55～59歳", "60～64歳", "65～69歳", "70～74歳", "中計")


class MhlwNdbParser(Parser):
    name = "mhlw-ndb"
    version = "2"
    source_kind = "mhlw_ndb"
    content_type = "*"

    def parse(self, snapshot: Snapshot, data: bytes, meta: Mapping[str, JsonValue] | None = None) -> ParseResult:
        if snapshot.sha256 is None:
            raise ParseError("NDB snapshot has no body")
        if meta and meta.get("member") != MEMBER:
            raise ParseError(f"NDB requires ZIP member {MEMBER!r}")
        members = read_zip(data)
        if MEMBER not in members:
            raise ParseError(f"Missing NDB ZIP member: {MEMBER}")
        payload = members[MEMBER]
        read_zip(payload)  # Bound and validate the nested XLSX before handing it to openpyxl.
        try:
            book = load_workbook(io.BytesIO(payload), read_only=True, data_only=False)
            try:
                if book.sheetnames != ["質問項目"]:
                    raise ValueError("expected the 質問項目 worksheet")
                sheet = book["質問項目"]
                sheet.reset_dimensions()  # Do not trust the producer's dimension declaration.
                rows = list(sheet.values)
            finally:
                book.close()
            if len(rows) != 101 or any(len(row) > 18 and any(v is not None for v in row[18:]) for row in rows):
                raise ValueError("unexpected NDB table dimensions")
            if (not isinstance(rows[0][0], str)
                    or not rows[0][0].startswith(f"特定健診 標準的な質問票（質問項目１） {INDICATOR}")
                    or "2023年度" not in rows[0][0]
                    or rows[1][:2] != ("都道府県名", "回答")
                    or rows[2][2] != "男" or rows[2][10] != "女"
                    or rows[3][2:18] != AGES * 2 or rows[4][2:18] != ("人数",) * 16):
                raise ValueError("unexpected NDB question, year or dimension headers")
            observations = []
            areas = set()
            for offset in range(5, 101, 2):
                area = rows[offset][0]
                if not isinstance(area, str) or not area or area in areas or rows[offset + 1][0] is not None:
                    raise ValueError("invalid or duplicate prefecture row pair")
                areas.add(area)
                for answer_offset, answer in enumerate(("はい", "いいえ")):
                    row = rows[offset + answer_offset]
                    if len(row) < 18 or row[1] != answer:
                        raise ValueError("unexpected answer row")
                    for col in range(2, 18):
                        raw = row[col]
                        note = None
                        if raw == "-":
                            value, note = None, "-; 遮蔽（集计单位小于 10 或防反推）"
                        elif raw is None or raw in ("…", "x", "秘匿"):
                            value, note = None, raw if raw is not None else "空欄"
                        elif type(raw) is int and raw >= 0:
                            value = Decimal(raw)
                        else:
                            raise ValueError(f"invalid count at row {offset + answer_offset + 1}, column {col + 1}: {raw!r}")
                        sex, age = ("男" if col < 10 else "女"), AGES[(col - 2) % 8]
                        cell = f"{get_column_letter(col + 1)}{offset + answer_offset + 1}"
                        observations.append(Observation(
                            obs_id=f"ndb11-q1:{area}:{sex}:{age}:{answer}", source_id=snapshot.source_id,
                            dataset_key=f"{snapshot.source_id}:ndb11-q1", table_key="ndb11-q1", survey=SURVEY,
                            indicator=INDICATOR, area_label=area, time_start=date(2023, 4, 1),
                            time_end=date(2024, 3, 31), time_label="2023年度", granularity=TimeGranularity.year,
                            dims={"sex": sex, "age": age, "answer": answer}, unit="人数", value=value, note=note,
                            snapshot_sha256=snapshot.sha256, fetched_at=snapshot.fetched_at,
                            release_id="ndb11", released_at=date(2026, 6, 16),
                            locator=format_locator({"member": MEMBER, "sheet": "質問項目", "cell": cell}),
                        ))
            # The A1 title carries the suppression rule; index it so it can be found and cited.
            dataset_key = f"{snapshot.source_id}:ndb11-q1"
            title_locator = format_locator({"member": MEMBER, "sheet": "質問項目", "cell": "A1"})
            document = Document(
                doc_key=dataset_key, source_id=snapshot.source_id, snapshot_sha256=snapshot.sha256,
                title=f"{SURVEY} {INDICATOR}", lang="ja", kind="dataset", content_type="application/zip",
            )
            title_chunk = Chunk(
                chunk_id=f"{dataset_key}#{title_locator}", doc_key=dataset_key, locator=title_locator,
                text=rows[0][0], title_path=[document.title], snapshot_sha256=snapshot.sha256, lang="ja",
            )
            return ParseResult(documents=[document], chunks=[title_chunk], observations=observations)
        except ParseError:
            raise
        except Exception as exc:
            raise ParseError(f"Invalid NDB questionnaire workbook: {exc}") from exc
