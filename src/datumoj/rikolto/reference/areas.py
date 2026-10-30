"""Lookup and normalize Japan's local government codes."""

from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook

_DATA_FILE = Path(__file__).with_name("data") / "000925835_dantai_code_R060101.xlsx"
_WEIGHTS = (6, 5, 4, 3, 2)


@dataclass(frozen=True, slots=True)
class Area:
    code5: str
    code6: str | None
    prefecture: str | None
    municipality: str | None
    prefecture_kana: str | None = None
    municipality_kana: str | None = None

    @property
    def label(self) -> str:
        if self.prefecture is None:
            return "全国"
        return self.prefecture + (self.municipality or "")


def normalize_area_label(s: str) -> str:
    """Normalize width and whitespace, joining prefecture and municipality labels."""
    return "".join(unicodedata.normalize("NFKC", s).split())


def code5_to_code6(code5: str) -> str:
    """Append the Ministry of Internal Affairs and Communications check digit."""
    code5 = normalize_area_label(code5)
    if not re.fullmatch(r"[0-9]{5}", code5):
        raise ValueError(f"expected a five-digit local government code: {code5!r}")
    total = sum(int(digit) * weight for digit, weight in zip(code5, _WEIGHTS))
    return code5 + str((11 - total % 11) % 10)


def code6_to_code5(code6: str) -> str:
    """Validate and remove a local government code's check digit."""
    code6 = normalize_area_label(code6)
    if not re.fullmatch(r"[0-9]{6}", code6):
        raise ValueError(f"expected a six-digit local government code: {code6!r}")
    code5 = code6[:5]
    if code6[-1] != code5_to_code6(code5)[-1]:
        raise ValueError(f"invalid local government code check digit: {code6!r}")
    return code5


@lru_cache(maxsize=1)
def _areas() -> tuple[Area, ...]:
    workbook = load_workbook(_DATA_FILE, read_only=True, data_only=True)
    try:
        by_code: dict[str, Area] = {}
        for sheet in workbook.worksheets[:2]:
            for row in sheet.iter_rows(min_row=2, values_only=True):
                code, prefecture, municipality, prefecture_kana, municipality_kana = row[:5]
                if not code or not prefecture:
                    continue
                code6 = str(code)
                code5 = code6_to_code5(code6)
                by_code.setdefault(
                    code6,
                    Area(
                        code5=code5,
                        code6=code6,
                        prefecture=str(prefecture),
                        municipality=str(municipality) if municipality else None,
                        prefecture_kana=str(prefecture_kana) if prefecture_kana else None,
                        municipality_kana=str(municipality_kana) if municipality_kana else None,
                    ),
                )
        return tuple(by_code[code] for code in sorted(by_code))
    finally:
        workbook.close()


_NATIONWIDE = Area(code5="00000", code6="000001", prefecture=None, municipality=None)


def _short_prefecture_name(name: str) -> str:
    return name[:-1] if name.endswith(("都", "府", "県")) else name


def lookup(name_or_code: str) -> tuple[Area, ...]:
    """Return all matching areas; ambiguous labels remain multiple candidates."""
    query = normalize_area_label(name_or_code)
    if query in {"00000", "000001", "全国"}:
        return (_NATIONWIDE,)

    areas = _areas()
    if query.isascii() and query.isdecimal():
        if len(query) == 2:
            return tuple(area for area in areas if area.municipality is None and area.code6[:2] == query)
        if len(query) == 5:
            code6 = code5_to_code6(query)
        elif len(query) == 6:
            try:
                code6_to_code5(query)
            except ValueError:
                return ()
            code6 = query
        else:
            return ()
        if code6 == _NATIONWIDE.code6:
            return (_NATIONWIDE,)
        return tuple(area for area in areas if area.code6 == code6)

    matches = []
    for area in areas:
        names = {area.prefecture, _short_prefecture_name(area.prefecture)}
        if area.municipality:
            names.update(
                {
                    area.municipality,
                    area.prefecture + area.municipality,
                    _short_prefecture_name(area.prefecture) + area.municipality,
                }
            )
        if area.prefecture_kana:
            names.add(area.prefecture_kana)
        if area.municipality_kana:
            names.add(area.municipality_kana)
            if area.prefecture_kana:
                names.add(area.prefecture_kana + area.municipality_kana)
        if query in {normalize_area_label(name) for name in names if name}:
            matches.append(area)
    return tuple(matches)
